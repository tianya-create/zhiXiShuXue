from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"D:\zhiXiShuXue")
OUT = ROOT / "output" / "reports"
ASSETS = ROOT / "tmp" / "report_assets"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / "智析数学研究报告_功能修订版.docx"

FONT = r"C:\Windows\Fonts\msyh.ttc"
NAVY = "173A78"
BLUE = "2577D7"
TEAL = "14A6A6"
PALE_BLUE = "EAF5FF"
PALE_TEAL = "E9F8F6"
PALE_YELLOW = "FFF6DE"
LINE = "D8E4F2"
GRAY = "53657E"

def cnfont(size, bold=False):
    return ImageFont.truetype(FONT, size, index=0)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=LINE):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = qn('w:' + edge)
        el = borders.find(tag)
        if el is None:
            el = OxmlElement('w:' + edge)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '6')
        el.set(qn('w:color'), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_font(run, size=10.5, bold=False, color=None, name='Microsoft YaHei'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def paragraph(doc, text='', size=10.5, bold=False, color=None, align=None, before=0, after=6, first_indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.55
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, 16 if level == 1 else 12.5, True, NAVY)
    return p

def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 9.5, False, GRAY)
    return p

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    h = table.rows[0]
    set_repeat_table_header(h)
    for i, value in enumerate(headers):
        cell = h.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(value)
        set_font(r, 9, True, 'FFFFFF')
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 0:
                set_cell_shading(cell, 'F6FAFE')
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            r = p.add_run(value)
            set_font(r, 8.7, False, '203657')
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def draw_round(draw, box, fill, outline=BLUE, radius=28, width=3):
    fill = fill if str(fill).startswith('#') else '#' + str(fill)
    outline = outline if str(outline).startswith('#') else '#' + str(outline)
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)

def center_text(draw, box, text, font, fill=NAVY, spacing=8):
    fill = fill if str(fill).startswith('#') else '#' + str(fill)
    left, top, right, bottom = box
    lines = text.split('\n')
    heights = [draw.textbbox((0,0), line, font=font)[3] for line in lines]
    total = sum(heights) + spacing*(len(lines)-1)
    y = top + (bottom-top-total)/2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0,0), line, font=font)[2]
        draw.text((left+(right-left-width)/2,y), line, font=font, fill=fill)
        y += height+spacing

def arrow(draw, x1,y1,x2,y2, fill=BLUE, width=8):
    fill = fill if str(fill).startswith('#') else '#' + str(fill)
    draw.line((x1,y1,x2,y2), fill=fill, width=width)
    # simple triangular head
    if x2 >= x1:
        pts=[(x2,y2),(x2-22,y2-12),(x2-22,y2+12)]
    else:
        pts=[(x2,y2),(x2+22,y2-12),(x2+22,y2+12)]
    draw.polygon(pts, fill=fill)

def make_flow(path):
    im=Image.new('RGB',(1800,900),'#FFFDF7'); d=ImageDraw.Draw(im)
    title=cnfont(52); label=cnfont(24); small=cnfont(20)
    d.text((90,55),'认知—心理双层适配闭环',font=title,fill='#173A78')
    boxes=[(95,265,390,500),(470,265,765,500),(845,265,1140,500),(1220,265,1515,500)]
    fills=['#EAF5FF','#FFF6DE','#E9F8F6','#EEF1FF']
    labels=['完成本轮练习\n记录正确率与任务复杂度','轮后自评\n控制感 + 心理努力','第一层决策\n任务复杂度、支架、信息密度','第二层呈现\n按资源偏好匹配展示方式']
    for b, f, txt in zip(boxes,fills,labels):
        draw_round(d,b,f)
        center_text(d,b,txt,label)
    for a,b in zip(boxes,boxes[1:]): arrow(d,a[2]+12,(a[1]+a[3])//2,b[0]-12,(b[1]+b[3])//2)
    draw_round(d,(420,630,1380,795),'#F7FBFF','#9BC8F5',32,3)
    center_text(d,(420,630,1380,795),'生成下一轮学习页面：题目与学习目标由第一层决定；图示、文字、例题或练习优先的资源呈现由第二层决定。',small)
    d.text((95,555),'适用范围：整式、整式加减、分式与分式运算、一元一次方程；其他知识点保持原有推送逻辑。',font=small,fill='#53657E')
    im.save(path)

def make_arch(path):
    im=Image.new('RGB',(1800,1030),'#FFFDF7'); d=ImageDraw.Draw(im)
    title=cnfont(52); label=cnfont(24); small=cnfont(19)
    d.text((90,50),'系统架构与数据边界',font=title,fill='#173A78')
    columns=[('学生端\n练习、自评、偏好测评、反馈', '#EAF5FF'),('教师端\n范围设置、审核、发布、统计', '#E9F8F6'),('管理端\n账号、规则、模型配置、备份', '#FFF6DE')]
    x=90
    for text, fill in columns:
        draw_round(d,(x,170,x+490,355),fill)
        center_text(d,(x,170,x+490,355),text,label)
        arrow(d,x+245,370,x+245,455)
        x+=560
    draw_round(d,(180,465,1620,690),'#EEF5FF','#2577D7',34,3)
    center_text(d,(180,465,1620,690),'服务层\n身份认证与角色校验｜作业、题库、审核与知识图谱\n自适应轮次与资源呈现｜AI 助教（OpenAI-compatible 服务，可配置）',label)
    arrow(d,900,705,900,780)
    draw_round(d,(330,795,1470,950),'#F8FAFC','#8DA7C5',32,3)
    center_text(d,(330,795,1470,950),'数据层：SQLite 持久化数据卷｜学习轮次、自评原始记录、适配决策、偏好记录\n上传文件与备份恢复。教育支持数据只服务于学习过程，不作心理健康诊断。',small)
    im.save(path)

def make_roles(path):
    im=Image.new('RGB',(1800,950),'#FFFDF7'); d=ImageDraw.Draw(im)
    title=cnfont(52); label=cnfont(31); small=cnfont(24)
    d.text((90,50),'师生协同使用流程',font=title,fill='#173A78')
    left=[('教师设置已学范围','#E9F8F6'),('题库审核与认知负荷标注','#EAF5FF'),('发布作业并查看班级分布','#FFF6DE')]
    right=[('首次资源呈现偏好测评','#EAF5FF'),('完成练习并提交轮后自评','#FFF6DE'),('查看反馈并进入下一轮','#E9F8F6')]
    for idx,(text,fill) in enumerate(left):
        y=190+idx*205; draw_round(d,(130,y,760,y+135),fill); center_text(d,(130,y,760,y+135),text,label)
        if idx<2: arrow(d,445,y+145,445,y+195,TEAL,7)
    for idx,(text,fill) in enumerate(right):
        y=190+idx*205; draw_round(d,(1040,y,1670,y+135),fill); center_text(d,(1040,y,1670,y+135),text,label)
        if idx<2: arrow(d,1355,y+145,1355,y+195,BLUE,7)
    draw_round(d,(795,370,1005,540),'#EEF1FF','#5576DD')
    center_text(d,(795,370,1005,540),'系统\n适配引擎',label)
    arrow(d,770,455,795,455,BLUE,7); arrow(d,1005,455,1030,455,TEAL,7)
    d.text((100,825),'教师查看的是学习支持记录和班级趋势；学生可主动发起“想和老师说”请求，AI 助教不替代教师的情感支持。',font=small,fill='#53657E')
    im.save(path)

make_flow(ASSETS/'flow.png')
make_arch(ASSETS/'architecture.png')
make_roles(ASSETS/'roles.png')

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.25); sec.bottom_margin = Cm(2.0); sec.left_margin = Cm(2.35); sec.right_margin = Cm(2.35)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)

# footer
footer = sec.footer
p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('智析数学研究报告 功能修订版'); set_font(r,8.5,False,'7B8CA8')

# cover
for _ in range(5): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('智析数学'); set_font(r,31,True,NAVY)
p.paragraph_format.space_after=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('初中数学智能学情分析网页系统'); set_font(r,21,True,NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('研究报告 功能修订版'); set_font(r,18,False,BLUE)
for _ in range(7): doc.add_paragraph()
for text in ['项目名称：智析数学', '修订主题：认知—心理双维度双层学情适配机制', '版本日期：2026 年 10 月']:
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(text);set_font(r,12,False,GRAY);p.paragraph_format.space_after=Pt(8)
doc.add_page_break()

heading(doc,'摘 要',1)
paragraph(doc,'智析数学是一套面向初中数学教学场景的智能学情分析网页系统。系统保留作业发布与完成、题库审核、自动批改、知识图谱、错题复习、学习轨迹、教师与管理员三端管理等原有能力，并在整式、整式加减、分式与分式运算、一元一次方程四类知识模块中新增认知—心理双维度双层学情适配机制。')
paragraph(doc,'该机制的第一层以本轮正答率、任务复杂度、情境性学业控制感与心理努力为共同依据，确定下一轮任务复杂度、支架强度、页面信息密度及情感引导语气；第二层读取学生的多模态学习资源呈现偏好，在不改变第一层确定的知识目标与任务难度的前提下，调整图示、文字、例题或练习优先的资源组织方式。学生每轮练习后完成自评，系统保存原始记录并生成可解释的学习反馈；教师端可查看班级分布与个体历史记录。')
paragraph(doc,'本报告将适配结果定位为学习支持证据，而非心理健康诊断或学生能力标签。AI 助教提供思路提示与学习陪伴，学生主动发起时可形成教师支持请求，但不替代教师的情感引导职责。', after=10)
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(12);r=p.add_run('关键词：');set_font(r,10.5,True,NAVY);r=p.add_run('初中数学；学情分析；认知负荷；学业控制感；资源呈现偏好；自适应学习');set_font(r,10.5,False,'203657')
heading(doc,'目录',1)
for item in ['第一章 研究背景与目标','第二章 双层适配机制设计','第三章 系统架构与功能实现','第四章 测试与验证设计','第五章 教育心理学创新与伦理边界','第六章 应用价值与后续计划']:
    p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(0.7);p.paragraph_format.space_after=Pt(5);r=p.add_run(item);set_font(r,11,False,NAVY)
doc.add_page_break()

heading(doc,'第一章 研究背景与目标',1)
heading(doc,'1.1 研究背景',2)
paragraph(doc,'传统数字化作业系统通常以得分、错题和知识点掌握度为主要输入，能够提高统计效率，却难以解释学生“为什么做不下去”或“为什么在会做时仍缺乏把握”。单纯依据正确率升降难度，可能把暂时的挫败、页面信息过密或支架不足误判为能力不足，也可能让高控制感学生停留在重复练习中。')
paragraph(doc,'为回应这一问题，项目在原有学情分析链路中引入“认知负荷—学业控制感”协同机制，并将资源偏好放在第二层作为呈现适配，而非独立决定任务难度的标签系统。')
heading(doc,'1.2 研究目标',2)
add_table(doc,['目标','具体内容'],[
['形成可解释的适配闭环','将答题表现、任务复杂度、学业控制感与心理努力整合为下一轮学习支持依据。'],
['维持教师主导','教师可设置已学范围、审核题目、查看班级与个体记录；AI 助教只提供辅助。'],
['提供差异化资源呈现','依据多模态资源呈现偏好组织图示、文字、例题或练习，保持同一学习目标。'],
['建立可追溯数据链','保存轮次、自评原始答案、适配决策及长期校正证据，支持复盘与测试。'],
], [4.0, 11.5])
heading(doc,'1.3 创新点与范围',2)
paragraph(doc,'本项目的核心创新不在于“自动升降难度”，而在于避免只凭正确率作机械判断：当低学业控制感与较高心理努力同时出现时，系统优先降低任务复杂度、拆分步骤、增加留白和支持性语气；当控制感较高且表现稳定时，系统提供同知识点变式或挑战，并使用明确的激励语气以避免自满。中等情形维持标准难度与布局。')
paragraph(doc,'双层适配仅在整式、整式加减、分式与分式运算、一元一次方程中启用。其他知识点沿用原有推送逻辑与页面样式，避免将尚未验证的规则扩展到全学科。')

heading(doc,'第二章 双层适配机制设计',1)
heading(doc,'2.1 核心闭环',2)
paragraph(doc,'学生完成一轮练习后，系统先计算正答率并读取本轮题目的认知负荷等级，再通过表情选择、两道学业控制感题和一道心理努力题获取自评。第一层输出任务复杂度、支架与页面密度；第二层再按照资源呈现偏好组装学习页面。两层按先后顺序运行，资源偏好不改变第一层确定的任务难度。')
doc.add_picture(str(ASSETS/'flow.png'), width=Cm(15.4)); caption(doc,'图 1 认知—心理双层适配闭环图')
doc.add_page_break()
heading(doc,'2.2 核心功能与实现要点',2)
add_table(doc,['功能','实现要点','数据与边界'],[
['轮后自评','表情组件、两道学业控制感题、一道心理努力题；不可空提交。','保存原始选择、得分与等级，仅用于学习支持。'],
['第一层适配','综合正答率、任务复杂度、控制感和心理努力，输出下一轮复杂度、支架、页面密度与情感语气。','低控制感优先拆分步骤；高控制感可进入变式或挑战。'],
['第二层适配','读取资源呈现偏好，加载图示、文字、例题或练习优先界面。','不改变知识目标、任务难度和教师已学范围。'],
['长期校正','累计至少 6 轮后评估；持续低效且存在接近的第二偏好时，扩展为双通道资源呈现。','证据不足时仅继续观察，不随意修改标签。'],
['情感支持','自评后按控制感结果出现可关闭便利贴；低控制感为安抚，高控制感为激励与要求。','不遮挡题目；学生可主动发起教师支持请求。'],
], [2.7, 7.0, 5.8])
heading(doc,'2.3 资源呈现偏好',2)
paragraph(doc,'首次登录时学生完成资源呈现偏好测评。系统将结果表述为“多模态学习资源呈现偏好”，而非稳定的学习能力或心理类型。视觉型优先呈现关系图和示意图；文字理解型以规范概念解析为主；例题示范型优先呈现分步例题；练习试错型保持简洁讲解并突出题目。')

heading(doc,'第三章 系统架构与功能实现',1)
heading(doc,'3.1 总体架构',2)
paragraph(doc,'系统采用 B/S 架构。前端负责角色界面、题目交互、自评与学习反馈呈现；后端负责身份认证、权限校验、作业与题库、知识图谱、适配规则和持久化数据服务。大模型能力采用可配置的 OpenAI-compatible 服务接口，能够对接合规的兼容服务；未配置时 AI 助教保持提示和降级状态，不影响作业、适配和统计主链路。')
doc.add_picture(str(ASSETS/'architecture.png'), width=Cm(15.4)); caption(doc,'图 2 系统架构与数据边界图')
heading(doc,'3.2 角色功能与师生协同',2)
paragraph(doc,'教师先设置已学范围并审核题目，题目在适配白名单模块中标注低、中、高认知负荷后可进入作业发布。学生完成作业、补练或复习后进入自评与反馈页。教师端新增控制感分布、题目认知负荷分布、资源呈现偏好分布，以及单个学生的历史自评和偏好记录查看能力。')
doc.add_picture(str(ASSETS/'roles.png'), width=Cm(13.5)); caption(doc,'图 3 师生协同使用流程图')
heading(doc,'3.3 技术栈与部署',2)
add_table(doc,['层级','当前实现','说明'],[
['前端','Vue 3、Vite、Vue Router、Element Plus、ECharts','学生、教师、管理员三端响应式界面；内部统一为笔记本风格。'],
['后端','Node.js、Express、REST API','处理认证、权限、作业、审核、统计与适配规则。'],
['数据','SQLite 持久化数据卷、上传文件持久化、备份恢复','保存学习轮次、自评原始数据、适配决策与偏好记录。'],
['模型服务','可配置 OpenAI-compatible 大模型服务','在管理员配置页设置 Endpoint、Model 与 API Key；密钥不在学生端暴露。'],
['部署','Docker Compose 或本地开发服务','生产环境需配置持久化卷、备份策略、HTTPS 与访问审计。'],
], [2.3, 5.1, 8.1])

heading(doc,'第四章 测试与验证设计',1)
heading(doc,'4.1 已完成的关键链路验证',2)
paragraph(doc,'当前开发版本已完成登录权限、教师范围设置、题目审核、作业发布、学生作答、自评提交、适配反馈、资源偏好长期校正、备份恢复及重启后数据持久化等核心链路检查。功能测试用于确认规则是否正确执行，不等同于教育效果已经得到实证验证。')
add_table(doc,['测试项','预期结果','验证重点'],[
['首次偏好测评拦截','未完成测评时进入测评页，完成后保存结果并加载相应资源呈现。','首次流程和持久化。'],
['自评提交校验','量表题或努力题为空时不能提交；提交后保存原始记录。','数据完整性。'],
['低控制感支持','在低控制感且投入较高时，下一轮降低复杂度、增强支架与留白。','第一层适配。'],
['高控制感挑战','在表现稳定且高控制感时，推送同知识点变式或挑战任务。','避免机械重复。'],
['白名单边界','非四类指定知识点保持原有推送逻辑。','生效范围。'],
['长期校正','满 6 轮后才评估；仅在持续低效且第二偏好接近时扩展双通道。','审慎校正。'],
['数据重启恢复','重启后轮次、自评、决策与偏好记录仍可查询。','持久化。'],
['教师数据权限','教师只能查看自己班级的学习支持记录。','权限与隐私。'],
], [3.0, 8.0, 4.5])
heading(doc,'4.2 生产环境收尾项',2)
paragraph(doc,'在正式部署前，仍需完成真实班级账户的端到端验收、浏览器兼容性复测、生产数据库备份演练、上传文件恢复演练、访问日志与敏感数据最小化检查。适配机制应以学习支持功能上线，后续通过对照研究评估其教育效果。')

heading(doc,'第五章 教育心理学创新与伦理边界',1)
heading(doc,'5.1 从正确率调节转向认知—心理协同支持',2)
paragraph(doc,'项目的教育心理学价值在于把学生的主观学习体验纳入任务调节依据。正答率反映表现结果，认知负荷反映任务要求，控制感反映学生对“能否推进”的判断，心理努力反映当前投入程度。四类信息共同使用，有助于避免将一时的困难简单归因为能力不足。')
heading(doc,'5.2 教师不可替代的情感引导',2)
paragraph(doc,'便利贴、表情和像素风数学老师形象承担的是即时、轻量、可关闭的情感支持：低控制感时提供安抚与分步提示，高控制感时提供激励和自我监控提醒。AI 助教不作心理诊断、不替代教师，也不以聊天结果改变任务难度；当学生需要更多支持时，主动发起的请求应由教师关注和回应。')
heading(doc,'5.3 数据伦理与解释边界',2)
paragraph(doc,'控制感、自评与资源偏好均为情境性学习支持数据，不应用于心理健康诊断、能力定型、升学评价或公开排名。教师端应以班级趋势与个体支持需要为主，限制跨班级访问；管理员需提供规则配置、备份恢复与访问控制。')

heading(doc,'第六章 应用价值与后续计划',1)
heading(doc,'6.1 应用价值',2)
paragraph(doc,'对教师而言，系统把作业、题库审核、已学范围、知识图谱与新增的适配统计纳入同一链路，使教师能够看到学生不仅“错在哪里”，也能看到“在什么状态下需要怎样的学习支持”。对学生而言，反馈页解释下一轮为何被拆分、为何展示例题或图示，降低黑箱式推荐带来的困惑。')
heading(doc,'6.2 后续教育效果验证',2)
paragraph(doc,'建议在真实教学中设置原推荐组、认知—心理适配组与双层适配组，在不影响常规教学的前提下比较迁移题表现、学习完成率、感知心理努力、控制感变化与教师支持请求处理情况。研究设计应明确样本量、年级、时间、测量来源和伦理告知，以避免用探索性数据作强结论。')
heading(doc,'6.3 报告图件更新建议',2)
add_table(doc,['原图号','处理建议','替换为'],[
['图 4','重画','图 1 认知—心理双层适配闭环图。'],
['图 5','拆分','图 2 系统架构与数据边界图；必要时另附技术部署图。'],
['图 6—图 8','更新','图 3 师生协同使用流程图，并补教师端适配统计截图。'],
['图 16—图 26','大部分替换','使用当前米白方格笔记本风格界面截图。'],
['图 23','必须替换','展示像素风数学老师形象的小智助教，而非旧 AI 图标。'],
['图 24、图 26','建议替换','轮后自评弹窗、便利贴支持与本轮学习反馈页。'],
], [2.0, 3.1, 10.4])
paragraph(doc,'注：正式答辩版应补充与当前系统一致的教师端统计、轮后自评、学习反馈页和资源呈现偏好页面截图，并以最新测试记录替换旧版单一回归测试图表。',size=9.5,color=GRAY,after=4)

heading(doc,'参考文献 新增与核对项',1)
references = [
    '[1] Sweller J. Cognitive load during problem solving Effects on learning[J]. Cognitive Science, 1988, 12(2): 257-285.',
    '[2] Paas F, Renkl A, Sweller J. Cognitive load theory and instructional design Recent developments[J]. Educational Psychologist, 2003, 38(1): 1-4.',
    '[3] Pekrun R. The control-value theory of achievement emotions Assumptions, corollaries, and implications for educational research and practice[J]. Educational Psychology Review, 2006, 18(4): 315-341.',
    '[4] Bandura A. Self-efficacy The exercise of control[M]. New York: W. H. Freeman, 1997.',
    '[5] Fleming N D, Mills C. Not another inventory Rather a catalyst for reflection[J]. To Improve the Academy, 1992, 11: 137-155.',
    '[6] Pashler H, McDaniel M, Rohrer D, Bjork R. Learning styles Concepts and evidence[J]. Psychological Science in the Public Interest, 2008, 9(3): 105-119.',
    '[7] 教育部. 义务教育数学课程标准 2022 年版[S]. 北京: 北京师范大学出版社, 2022.',
    '[8] 国家互联网信息办公室. 生成式人工智能服务管理暂行办法[Z]. 2023.',
]
for ref in references:
    paragraph(doc, ref, size=9.5, color='203657', after=4, first_indent=False)
paragraph(doc,'说明：正式合并时请保留原报告中已经核验的政策、教育信息化、软件测试和数据保护文献，并将以上新增条目与对应章节的文内引文一一核对。',size=9.2,color=GRAY,after=4,first_indent=False)

doc.save(OUTPUT)
print(OUTPUT)
